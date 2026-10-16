import hashlib
from pathlib import Path
import sys
import tempfile
import unittest
ROOT=Path(__file__).resolve().parents[1]
sys.path.insert(0,str(ROOT))
from backend import local_files as files

class LocalFiles(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.path=Path(self.temp.name)
    def tearDown(self):self.temp.cleanup()
    def write(self,name,text,encoding='utf-8'):
        path=self.path/name;path.write_text(text,encoding=encoding);return path
    def test_read_code_unchanged_and_list_subfolders(self):
        path=self.write('模块.py','print("读取SENTINEL")')
        (self.path/'子目录').mkdir()
        before=hashlib.sha256(path.read_bytes()).hexdigest()
        self.assertIn('SENTINEL',files.document(path)['content'])
        self.assertEqual(hashlib.sha256(path.read_bytes()).hexdigest(),before)
        self.assertEqual(files.listing(self.path)['entries'][0]['name'],'子目录')
    def test_long_file_pages_have_explicit_ranges(self):
        path=self.write('long.log','A'*40000+'TAIL')
        first=files.document(path);second=files.document(path,first['end'])
        self.assertTrue(first['partial']);self.assertEqual(first['end'],39000)
        self.assertEqual(first['content']+second['content'],'A'*40000+'TAIL')
    def test_credentials_not_available_for_reference(self):
        path=self.write('auth.json','PRIVATE_SENTINEL')
        doc=files.document(path)
        self.assertFalse(doc['attachable']);self.assertNotIn('PRIVATE_SENTINEL',doc['content'])
    def test_noncredential_secret_field_is_redacted(self):
        path=self.write('sample.json','{"apiKey":"VALUEVERYSECRET"}')
        self.assertNotIn('VALUEVERYSECRET',files.document(path)['content'])
    def test_binary_and_images_have_explicit_unsupported_status(self):
        path=self.path/'opaque.bin';path.write_bytes(bytes(range(256)))
        self.assertFalse(files.document(path)['attachable'])
        image=self.path/'image.png';image.write_bytes(b'fake')
        self.assertIn('图片',files.document(image)['reason'])
    def test_gb18030_and_utf16_text(self):
        for encoding in ('gb18030','utf-16'):
            self.assertEqual(files.document(self.write(encoding+'.txt','中文资料',encoding))['content'],'中文资料')
    def test_search_names_and_bounded_subfolders(self):
        (self.path/'sub').mkdir();self.write('noise.txt','N')
        (self.path/'sub/target.md').write_text('T')
        self.assertEqual(files.listing(self.path,'target')['entries'],[])
        self.assertEqual(len(files.listing(self.path,'target',True)['entries']),1)
    def test_relative_and_unc_paths_rejected(self):
        for value in ('relative.txt','\\\\remote\\file.txt','//remote/file.txt'):
            with self.assertRaises(ValueError):files.absolute_path(value)
    def test_docx_and_excel_extraction_does_not_modify(self):
        from docx import Document
        from openpyxl import Workbook
        doc=Document();doc.add_paragraph('WORD_SENTINEL');word=self.path/'test.docx';doc.save(word)
        book=Workbook();book.active['A1']='SHEET_SENTINEL';excel=self.path/'test.xlsx';book.save(excel);book.close()
        for path,sentinel in ((word,'WORD_SENTINEL'),(excel,'SHEET_SENTINEL')):
            before=path.read_bytes();self.assertIn(sentinel,files.document(path)['content']);self.assertEqual(before,path.read_bytes())

if __name__=='__main__':unittest.main()
