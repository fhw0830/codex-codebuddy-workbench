"""User-selected local browsing and bounded, read-only document extraction."""
import hashlib
import os
import re
from pathlib import Path
import time
import zipfile
from .store import clean

TEXT_EXTENSIONS=set('.md .txt .log .py .js .jsx .ts .tsx .html .css .scss .json .jsonl .yaml .yml .toml .ini .cfg .xml .csv .sql .sh .ps1 .bat .cmd .c .cpp .cc .h .hpp .java .kt .go .rs .vue .svelte .properties .srv .msg .action .launch .cmake .tex .rst .gitignore'.split())
IMAGES={'.png','.jpg','.jpeg','.gif','.webp','.bmp'}
MAX_BYTES=20*1024*1024
MAX_CHARS=39000

def absolute_path(value):
    value=str(value)
    if value.startswith(('\\\\','//')) or '\x00' in value:
        raise ValueError('请选择本机磁盘上的文件夹或文件')
    path=Path(value).expanduser()
    if not path.is_absolute():
        raise ValueError('请填写完整的本地路径')
    path=path.resolve()
    if os.name=='nt' and ':' in str(path)[2:]:
        raise ValueError('不支持设备路径或备用数据流')
    return path

def protected(path):
    name=path.name.lower()
    return '.ssh' in {p.lower() for p in path.parts} or name in {'auth.json','credentials','credentials.json','models.json','config.local.json','id_rsa','id_ed25519'} or name.startswith('.env') or path.suffix.lower() in {'.pem','.key','.pfx','.p12','.kdbx'}

def roots():
    if os.name=='nt':
        import ctypes
        mask=ctypes.windll.kernel32.GetLogicalDrives()
        return [{'name':chr(65+i)+':','path':chr(65+i)+':/','directory':True} for i in range(26) if mask & (1<<i)]
    return [{'name':'本机文件','path':'/','directory':True}]

def listing(value='',q='',recursive=False):
    if not value:
        return {'path':'','parent':None,'entries':roots(),'limited':False}
    path=absolute_path(value)
    if not path.is_dir():
        raise ValueError('文件夹不存在或无法读取')
    needle=str(q).lower().strip()
    entries=[];examined=0;limited=False;until=time.monotonic()+3;pending=[(path,0)]
    while pending:
        folder,depth=pending.pop()
        try:
            with os.scandir(folder) as iterator:
                for item in iterator:
                    examined+=1
                    if examined>10000 or time.monotonic()>until:
                        limited=True;break
                    try:
                        directory=item.is_dir(follow_symlinks=False)
                        if recursive and needle and directory and depth<3 and not item.is_symlink():
                            pending.append((Path(item.path),depth+1))
                        if needle and needle not in str(Path(item.path).relative_to(path)).lower():
                            continue
                        stat=item.stat(follow_symlinks=False)
                        entries.append({'name':item.name,'path':Path(item.path).as_posix(),'directory':directory,
                            'bytes':stat.st_size,'modified':stat.st_mtime,'protected':protected(Path(item.path))})
                        if len(entries)>=500:
                            limited=True;break
                    except OSError:
                        continue
        except OSError:
            if folder==path:
                raise ValueError('该文件夹没有读取权限')
        if limited:break
    return {'path':path.as_posix(),'parent':path.parent.as_posix() if path.parent!=path else '',
        'entries':sorted(entries,key=lambda e:(not e['directory'],e['name'].lower())),
        'limited':limited,'examined':examined,'search_depth':4 if recursive else 1}

def zip_check(path):
    with zipfile.ZipFile(path) as archive:
        if sum(item.file_size for item in archive.infolist())>32*1024*1024:
            raise ValueError('文档展开后过大，请提供相关片段')

def document(value,start=0,limit=MAX_CHARS):
    path=absolute_path(value)
    if not path.is_file():raise ValueError('文件不存在或不是普通文件')
    stat=path.stat()
    result={'local_path':path.as_posix(),'name':path.name,'bytes':stat.st_size,'content':'',
        'attachable':False,'format':path.suffix.lower().lstrip('.') or '无扩展名'}
    if protected(path):return {**result,'reason':'账号配置、私钥或凭据文件不提供内容引用。'}
    if stat.st_size>MAX_BYTES:return {**result,'reason':'文件超过20 MB，请选取相关日志或较小的文档片段。'}
    suffix=path.suffix.lower()
    try:
        if suffix=='.pdf':
            from pypdf import PdfReader
            reader=PdfReader(path)
            text='\n\n'.join(f'[第{i+1}页]\n'+(page.extract_text() or '') for i,page in enumerate(reader.pages[:50]))
            result['extraction_note']='最多前50页文本；扫描图片不自动OCR。'
        elif suffix=='.docx':
            zip_check(path)
            from docx import Document
            doc=Document(path)
            text='\n'.join(p.text for p in doc.paragraphs)
            for table in doc.tables:
                text+='\n'+'\n'.join('\t'.join(c.text for c in row.cells) for row in table.rows)
        elif suffix in {'.xlsx','.xlsm'}:
            zip_check(path)
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=True);blocks=[]
            try:
                for sheet in book.worksheets[:10]:
                    blocks.append('[工作表：'+sheet.title+']')
                    blocks.extend('\t'.join('' if v is None else str(v) for v in row) for row in sheet.iter_rows(max_row=300,max_col=30,values_only=True))
            finally:book.close()
            text='\n'.join(blocks)
            result['extraction_note']='最多10张表，每表300行/30列；公式使用文件内已有缓存值，未重算。'
        elif suffix=='.pptx':
            zip_check(path)
            import xml.etree.ElementTree as ET
            with zipfile.ZipFile(path) as archive:
                slides=sorted((n for n in archive.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)),key=lambda n:int(Path(n).stem[5:]))[:50]
                text='\n\n'.join('[幻灯片：'+n+']\n'+'\n'.join(e.text or '' for e in ET.fromstring(archive.read(n)).iter() if e.tag.endswith('}t')) for n in slides)
        elif suffix in IMAGES:
            return {**result,'reason':'图片可直接拖入对话输入框，或用附件按钮添加后交给Codex查看。此入口只提取文档文本。'}
        else:
            blob=path.read_bytes()
            if suffix not in TEXT_EXTENSIONS and (b'\x00' in blob[:8000] or any(v<9 for v in blob[:1000])):
                return {**result,'reason':'该二进制格式暂不能提取文本；可以查看文件信息。'}
            text=None
            encodings=('utf-16','utf-8-sig','gb18030') if blob.startswith((b'\xff\xfe',b'\xfe\xff')) else ('utf-8-sig','gb18030')
            for encoding in encodings:
                try:
                    candidate=blob.decode(encoding)
                    if '\x00' not in candidate[:8000]:text=candidate;break
                except UnicodeError:continue
            if text is None:return {**result,'reason':'无法识别编码，请先转换为UTF-8文本。'}
        text=clean(text);start=max(0,int(start));limit=min(MAX_CHARS,max(1,int(limit)));excerpt=text[start:start+limit]
        return {**result,'content':excerpt,'attachable':bool(excerpt.strip()),'start':start,'end':start+len(excerpt),
            'total_chars':len(text),'partial':start>0 or start+len(excerpt)<len(text),
            'sha256':hashlib.sha256(excerpt.encode('utf-8')).hexdigest(),
            'reason':'' if excerpt.strip() else '未提取到文本；可能是扫描文档或选择范围为空。'}
    except ImportError:return {**result,'reason':'当前Python缺少该格式的读取库，请使用文本附件。'}
    except Exception as exc:raise ValueError('无法提取该文件，请检查权限或文档格式') from exc
